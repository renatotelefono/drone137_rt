"""Build the ground-computer design document with the bundled document runtime."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BOM = json.loads((ROOT / 'hardware/ground-computer-bom.json').read_text())
OUT = ROOT / 'hardware/drone137-ground-computer-project.docx'
doc = Document()
# The bundled Word template contains title/heading borders; this proposal uses
# whitespace and typography rather than those inherited decorative rules.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.72)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 26), ('Heading 1', 17), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(12)
    s.paragraph_format.space_after = Pt(8)
doc.core_properties.title = 'Low cost drone137 aircraft project'
doc.core_properties.subject = 'Ground computer AI and a lightweight onboard video and flight interface'
doc.core_properties.author = 'drone137'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break(); doc.add_heading(title, 1)

def table(headers, rows, widths, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if ri==0 or ri%2==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244153' if ri==0 else 'F2F5F7');pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:
                    run.font.size=Pt(size)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    p('')
    return t

def linkpara(label, url, note=''):
    para=p('')
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),rel)
    run=OxmlElement('w:r');rpr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245875');rpr.append(color);run.append(rpr)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyp.append(run);para._p.append(hyp)
    if note:para.add_run(' — '+note)
    return para

from decimal import Decimal

def total(items, key):
    return sum(Decimal(str(x['quantity'])) * Decimal(str(x[key])) for x in items)

def money(value):
    return f'${value:,.2f}'

air = total(BOM['aircraft_items'], 'prototype_unit_usd')
target = total(BOM['aircraft_items'], 'target_unit_usd')
ground = total(BOM['shared_ground_items'], 'unit_budget')
spares = total(BOM['initial_spares'], 'unit_budget')
mass = total(BOM['aircraft_items'], 'unit_mass_g')
assert (air, target, ground, spares, mass) == tuple(map(Decimal, ['256.48','184.99','181.5','22','216.3']))
refs = {'fc_store':1,'fc_manual':2,'frame':3,'frame_price':4,'motor':5,'motor_price':6,'propulsion':7,'pi':8,'camera':9,'camera_software':10,'flow_store':11,'flow_manual':12,'offboard':13,'radio':14,'gnss':15}

doc.add_heading('drone137\nCost-first micro quadcopter', 0)
p('Ground computer AI | Lightweight aircraft | Hardware project, revision 2')
p('21 October 2026 | Prices and budgets in USD')
p('Design objective: aircraft parts below $200, with all mission planning, recognition and reporting on a ground computer with an 8 GB RAM target. Use a small 3.5-inch quadcopter, an integrated PX4 flight controller/ESC board, a camera and a Pi Zero 2 W that only encodes video and forwards telemetry.')
table(['Design target','Current status'],[
('Aircraft parts below $200',f'{money(target)} procurement goal; {money(air)} prototype allowance. Lower prices are unquoted; the cost target is not yet met.'),
('Lightweight with battery and guards','230–280 g provisional planning range. The component budget plus 15% mass margin is about 249 g; no weighed assembly exists.'),
('Computer-controlled inspection','Ground computer plans bounded moves. PX4 stabilizes the aircraft and handles configured failsafes.'),
('Small controlled test area first','Prepared indoor bay with visible localization markers, adequate floor texture and lighting. Unfamiliar-room autonomy needs further sensing and validation.')],[2,4.8],10)
h('Where each function runs')
table(['Location','Responsibility'],[
('8 GB ground computer','drone137, local models, deterministic flight executor, video decoding, localization, mission reporting and recordings.'),
('PX4 autopilot onboard','IMU/barometer estimation, attitude control, motor mixing, battery monitoring and commissioned link-loss behaviour.'),
('Pi Zero 2 W onboard','Hardware H.264 encoding and MAVLink network/serial routing. It runs no LLM or object-recognition model.')],[1.7,5.1],10)
p('Command path: ground executor → local Wi-Fi → Pi UART bridge → PX4 → integrated ESCs → four motors. Camera video and telemetry return to the computer. A separate RC receiver connects directly to PX4 for human takeover.')
p('The first aircraft is an engineering prototype. Component references establish a feasible direction, but mounts, exact suppliers, sensor integration, thrust, flight duration and radio coverage still require validation. This document does not enable physical flight in drone137.')

page('Aircraft bill of materials')
p('Prototype figures combine checked listings and explicit allowances. Target figures are maximum procurement prices to seek for repeat builds, not supplier quotations or promised volume discounts. Both columns include one battery, guards and the optical-flow/range sensor; shared ground equipment is separate.')
notes = {
'frame':'Allowance; Crux35 reference $24.99, out of stock [3,4]',
'guards':'Allowance; mount/clearance design required',
'fc':'Listed $64.99 [1]; $50 target is unquoted',
'motors':'Allowance $12 each; $11.99 reference out of stock [5,6]',
'props':'Allowance; matched family reference [7]',
'pi':'Manufacturer starting price [8]',
'camera':'Listed sale price [9]; confirm package',
'flow':'Listed $16.50 [11]; $14 target unquoted',
'battery':'Allowance; exact pack and current rating unselected',
'receiver':'Allowance; CRSF serial output required',
'regulator':'Allowance; 4S-compatible input required',
'storage':'Allowance; two separate cards',
'ffc':'Allowance; confirm camera/Zero connectors',
'hardware':'Allowance including mounts and landing feet',
'buzzer':'Allowance'
}
rows=[]
for x in BOM['aircraft_items']:
    rows.append((x['item'], str(x['quantity']), f"{x['quantity']*x['prototype_unit_usd']:.2f}", f"{x['quantity']*x['target_unit_usd']:.2f}", notes[x['id']]))
rows.append(('Aircraft parts total','',f'{air:.2f}',f'{target:.2f}','Tax, shipping, labour and contingency excluded'))
table(['Component','Qty','Prototype\nUSD','Target\nUSD','Evidence / remaining selection'],rows,[2.18,.35,.69,.64,2.94],9)
h('What must change to reach the target')
p(f'The prototype allowance exceeds $200 by {money(air-Decimal(200))}. The target seeks {money(air-target)} of reductions: $20 across the four motors, $14.99 on the AIO board, $10 on the frame and $26.50 across the remaining items. These reductions require compatible supplier offers and a validated repeat build. Do not assume a discount from a particular batch size.')
p(f'The {money(target)} goal leaves {money(Decimal(200)-target)} before the $200 ceiling. A 5% parts reserve gives {money(target*Decimal("1.05"))}, still excluding delivery, tax and labour. An unavailable listing is a reference price, not a ready-to-order offer. If compatible quotes exceed $200, record the miss and retain the true cost.')

page('Mechanical, power and wiring design')
h('Small quad-X configuration')
p('Use a 150 mm motor-diagonal, 3.5-inch frame as the reference geometry. Happymodel documents EX1404 3500KV motors, HQProp T3.5X2X3 propellers and 4S power in the Crux35 family [3,5,7]. Select a 650–750 mAh pack for this heavier camera-carrying build; the stock drone’s flight time does not transfer to it. Four motors need two CW and two CCW propellers.')
p('The frame reference has M2 mounts; the AIO uses 25.5 mm spacing and 3 mm holes [2,3]. Confirm the rotated board envelope, insulated grommet sleeves and USB access in a mounting drawing. The 65 × 30 mm Pi [8] and 25 × 24 mm camera [9] need a separate mount; the stock 14 mm FPV-camera cage is unsuitable. Keep the antenna and ESC cooling area clear.')
p('Guards are included as an allowance, not a validated part. Check propeller clearance under deflection, landing clearance for the flow sensor and camera visibility. Guards reduce incidental contact but do not establish collision tolerance. Measure the assembled centre of gravity, inertia and actual rotor locations before flight modelling.')
table(['Mass budget','Grams','Basis'],[
('Frame and guards','46','26 g reference frame + 20 g guard allowance [3]'),
('Four motors and propeller set','38.8','4 × 8.7 g motors + 4 g prop allowance [5]'),
('Flight controller / ESC board','10','Manufacturer board mass [2]'),
('Pi, camera, regulator, cards and ribbon','23','Engineering allowances'),
('Receiver, flow sensor and buzzer','5.5','1.5 g flow sensor [12]; other allowances'),
('Battery / mounting hardware and wiring','75 / 18','Engineering allowances'),
('Nominal total / with 15% mass margin',f'{mass:.1f} / {mass*Decimal("1.15"):.1f}','Not measured; a sub-250 g result is not guaranteed')],[2.65,.85,3.3],9)
table(['Connection','Proposed route'],[
('Power','4S battery → AIO VBAT; separate regulated 5V branch → Pi. Common ground, correct polarity and rated capacitor/wiring.'),
('Video','B0390 CSI camera → compatible Zero ribbon → Pi. Verify pin count/contact orientation; strain-relieve the cable.'),
('MAVLink','Pi 3.3V UART ↔ FC UART1 / TELEM1; TX/RX crossed.'),
('Flow / manual receiver','MTF-02P → UART4 / TELEM2; CRSF receiver → UART6 / RC. Use rated 5V power and verify the exact pinout [2].')],[1.05,5.75],9)
p('Use the exact MicoAir743v2-AIO-35A board and its documented PX4 target, micoair_h743-v2_default [2]. Freeze firmware and parameters after bench checks. The AIO’s 12V output must not feed the Pi or 5V sensors. An oversized ESC current rating does not prove adequate motor, wiring or battery performance.')

page('drone137 integration and positioning')
h('A workable path from simulator to aircraft')
table(['Existing repo behaviour','Required integration'],[
('Loopback SITL only in adapters/mavsdk_drone.py and app/native.py','Add a separate physical launcher and explicit vehicle/network selection. Validate actual home/local origin, identity, telemetry health and completed actions.'),
('Loopback H.264 receiver; 320 × 240 at 2 fps and receipt timestamps','Add physical video input, configurable resolution, frame IDs and capture-clock mapping. Keep recording quality separate from model input size.'),
('Native flights use synthetic odometry and vertical-column range','Replace both with measured sensors. A downward sensor cannot supply the simulator’s wall/obstacle observations.'),
('SQLite profiles enforce large-airframe limits','Add a genuine micro-quad geometry and supported parameter ranges; 3.5-inch rotor radius is 0.04445 m, below the present 0.06 m minimum.'),
('app/contracts.py requires physical_enabled=False','Retain that gate while integrating and testing. Enable a real-vehicle mode only after explicit commissioning and acceptance.')],[2.65,4.15],9.5)
p('Do not shrink the rendered x500 and call it a calibrated small drone. The current minimum mass is 0.25 kg and minimum inertia is 0.001 kg·m². Use measured small-airframe values, independently placed rotors, a guard collision envelope and a calibrated motor/prop curve. Store revisions in SQLite so different aircraft can be selected without overwriting earlier profiles.')
h('Position is the remaining indoor dependency')
p('MTF-02P supplies downward range and optical flow, with manufacturer requirements including floor texture and more than 60 lux [12]. It can support local velocity and height estimation; integrated position drifts. It does not see walls or identify a global location in a room. Its advertised range is conditional on light and reflectance; measure it on the actual floor.')
p('For the first prepared bay, place measured visual markers where the forward camera can observe them. The ground computer must estimate camera pose, transform it to the FC/body frame, account for delay and send validated external-vision measurements to PX4. Marker detection, map coordinates, clock synchronization and estimator fusion are development work, not existing capability. Loss of valid positioning must stop mission progression.')
p('Use a known, cleared operating volume for initial movement. Unfamiliar workshop/warehouse search requires validated obstacle sensing and localization beyond this base bill of materials. Outdoor perimeter/field work can add GNSS plus a compass and mount, with a $23 allowance [15]; it still needs a surveyed operating area, link coverage and obstacle clearance. Hillside terrain following is not qualified by one downward range sensor.')
h('Model and flight-control separation')
p('A deterministic executor streams bounded setpoints independently of inference. PX4 requires offboard proof of life above 2 Hz [13]; 20 Hz is the proposed test setting. The Pi must not replay a stale command indefinitely after the ground process stops. Test loss actions on PX4; an LLM is not the attitude controller.')

page('Commissioning, memory and full cost')
h('Validation sequence and pass criteria')
p('1. Close the design. Select exact SKUs, resolve board/camera/guard mounts and connectors, then update the cost and mass sheet. Accept the under-$200 claim only when the complete quoted aircraft BOM meets it. Confirm full-pack and low-battery thrust margin on a restrained thrust stand; a proposed requirement is at least 2:1 maximum thrust to takeoff weight near the lower operating voltage.')
p('2. Bench with propellers removed. Verify power, motor output mapping and directions, RC takeover, camera calibration, clocks, FC logging and flow orientation. Run simultaneous video and telemetry for 30 minutes. Record frame age, round-trip delay, lost packets, temperatures and memory. Initial engineering targets: p95 frame age below 250 ms and command round-trip below 100 ms; these are not measured results.')
p('3. Inject failures before autonomous tests. Stop the model, executor, camera, Wi-Fi and Pi separately. Check stale-frame handling and configured PX4 loss actions. A return command needs valid position and a clear route; the prepared indoor bay needs an appropriate tested hold/land policy. Do not mask command loss by continuously resending the last move.')
p('4. Commission the real aircraft. A human operator verifies lift-off, hover, landing and takeover in a clear test area. Start at 0.3 m/s. Measure tracking error, camera blur, voltage sag, thrust and landing reserve; no endurance claim is made before these measurements. Replace provisional simulation parameters with measured profiles.')
p('5. Qualify missions on an actual 8 GB computer. Run target-present and target-absent trials with varied viewpoints and lighting. Retain failed flights and false confirmations. Require every planned qualification case to pass; a finite all-pass suite still does not establish universal accuracy. The existing Qwen + Florence development result is 3/4, and therefore does not meet that requirement.')
h('Keep the ground application light')
p('Start with one aircraft and one shared model service. Use one newest-frame slot per aircraft and bounded inference queues. Record H.264 without decoding every frame; decode sampled observations and detailed crops for inspection. Begin with 720p at 15 fps / 2–4 Mbit/s as link-test settings. Use the Pi hardware encoder [8,10], keep the 3D simulator off during physical operation and measure total host RAM, swap and response time.')
p('Stop and inspect at observation points so model latency does not affect stabilization. Qwen + Florence is the strongest tested pair so far; 4.47 GiB peak application memory was measured in the original comparison on a 16 GiB host. It does not establish 8 GB performance or a supported multi-drone count. Fleet capacity must be measured with aggregate video traffic and concurrent failure handling.')
h('Shared equipment and cost to produce')
table(['Budget item','USD'],[
('Shared: ELRS transmitter [14] + cells + 4S charger/PSU + Wi-Fi AP + marker materials','71.50 + 10 + 60 + 30 + 10'),
('Initial spares: second battery + propeller set',f'{spares:.2f}'),
('First prototype system, excluding existing computer',f'{air+ground+spares:.2f}'),
('Same prototype system with 10% allowance',f'{(air+ground+spares)*Decimal("1.1"):.2f}')],[4.7,2.1],9)
p(f'Parts are not finished manufacturing cost. At the unquoted {money(target)} parts goal plus an assumed 1.5–3 hours of assembly/test at $20/hour, direct cost would be about $215–245 before tax, shipping, engineering, tooling, overhead and rejects. Standardize one harness and mount set after validation; price batch builds only against actual supplier quotes.')

page('Sources and evidence')
p('Listings checked 21 October 2026. Prices can change; allowances and procurement targets are engineering inputs. Confirm exact variants, availability and complete delivered costs before ordering. Sources below are clickable.')
sources=[
('fc_store','[1] MicoAir743v2-AIO-35A store','Listed $64.99, integrated FC and four ESCs.'),
('fc_manual','[2] MicoAir AIO board manual','Dimensions, mass, UARTs, power and PX4 target.'),
('frame','[3] Happymodel Crux35 frame','Reference geometry, mounting and 26 g frame.'),
('frame_price','[4] GetFPV Crux35 frame listing','Reference $24.99; out of stock when checked.'),
('motor','[5] Happymodel EX1404 motors','3500KV configuration and 8.7 g motor mass.'),
('motor_price','[6] Fractal Engineering motor listing','Reference $11.99; out of stock when checked.'),
('propulsion','[7] Happymodel Crux35 ELRS V2 manual','EX1404 / T3.5X2X3 / 4S propulsion reference.'),
('pi','[8] Raspberry Pi Zero 2 W','Starting $15; H.264 encoder and 2.4 GHz Wi-Fi.'),
('camera','[9] Arducam B0390 IMX219','Listed $9.99 sale price; fixed focus visible-light camera.'),
('camera_software','[10] Raspberry Pi camera software','rpicam/libcamera capture and streaming.'),
('flow_store','[11] MicoAir MTF-02P store','Listed $16.50.'),
('flow_manual','[12] MicoAir MTF-02P specifications','Optical flow, conditional range, 1.5 g and PX4 support.'),
('offboard','[13] PX4 offboard documentation','Setpoint stream, estimator prerequisites and loss actions.'),
('radio','[14] RadioMaster Pocket','ELRS variant $71.50; cells excluded.'),
('gnss','[15] MicoAir M9/M10 GNSS family','$13.99–17.99 family pricing; variant selection pending.')]
for key,label,note in sources:linkpara(label,BOM['sources'][key],note)
h('Repository evidence and deliverables')
p('Integration review: adapters/mavsdk_drone.py, adapters/camera.py, adapters/native_flight.py, app/native.py, app/contracts.py, app/airframes.py and sim/hardware.py. Existing simulation and physical-mode restrictions remain in place. The micro-quad model, marker pose pipeline and physical transport are required work, not completed features.')
p('Model evidence: docs/vision-flight-comparison.md and docs/vision-flight-validation.json. The known synthetic development flights use simulated sensor inputs. Results do not certify this proposed aircraft, real-environment recognition or an 8 GB host.')
p('ground-computer-bom.json stores quantities, prototype allowances, procurement ceilings, mass assumptions, source URLs and exclusions. build_ground_computer_project.py reproduces this document and checks its totals. No purchases, assembly, bench measurements or real flights were performed for this design.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(json.dumps({'output':str(OUT),'prototype_aircraft_usd':str(air),'target_aircraft_usd':str(target),'target_is_quoted':False,'shared_ground_usd':str(ground),'spares_usd':str(spares),'first_system_usd':str(air+ground+spares),'mass_budget_g':str(mass)}))
