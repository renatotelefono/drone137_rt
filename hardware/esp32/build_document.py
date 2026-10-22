"""Build the ground-computer design document with the bundled document runtime."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BOM = json.loads((ROOT / 'hardware/esp32/bom.json').read_text())
OUT = ROOT / 'hardware/esp32/drone137-esp32-build.docx'
doc = Document()
# The bundled Word template contains title/heading borders; this proposal uses
# whitespace and typography rather than those inherited decorative rules.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.72)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 26), ('Heading 1', 17), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(12)
    s.paragraph_format.space_after = Pt(8)
doc.core_properties.title = 'drone137 ESP32 vision drone build'
doc.core_properties.subject = 'Ground computer AI and a lightweight onboard video and flight interface'
doc.core_properties.author = 'drone137'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break(); doc.add_heading(title, 1)

def table(headers, rows, widths, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if ri==0 or ri%2==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244153' if ri==0 else 'F2F5F7');pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:
                    run.font.size=Pt(size)
                    if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    p('')
    return t

def linkpara(label, url, note=''):
    para=p('')
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),rel)
    run=OxmlElement('w:r');rpr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245875');rpr.append(color);run.append(rpr)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyp.append(run);para._p.append(hyp)
    if note:para.add_run(' — '+note)
    return para

from decimal import Decimal
basic=Decimal(BOM['basic_estimate_usd'])
nav=Decimal(BOM['navigation_estimate_usd'])
assert basic==Decimal('48.50') and nav==Decimal('69.00')
doc.add_heading('drone137 ESP32 vision drone build',0)
p('Revision A   22 October 2026   Ground computer AI')
p('Build a lightweight quadcopter that sends live camera images to the drone137 harness over local Wi-Fi. The ground computer runs the models, records evidence and plans bounded movement. Stabilization and command-loss handling belong on the aircraft.')
p('This package defines the pieces, wiring, frame geometry and camera path. The basic parts allowance is $48.50; the version with optical flow, downward range and its power supply is $69.00. These are engineering allowances with two checked component listings, not a fully quoted purchase. The under-$50 objective is not yet verified.')
table(['Element','Selected direction'],[
('Aircraft','100 mm motor-diagonal quad-X, four 8520 brushed motors, 55 mm propellers and guards; target assembled mass of 60 to 85 g.'),
('Electronics','XIAO ESP32-S3 Sense with supplied camera, MPU6050 IMU and four low-side MOSFET motor drivers.'),
('Ground computer','Local drone137 application and model service; all recognition, mission planning and recording remain here.'),
('Vision transport','640 × 480 JPEG over HTTP; PC samples the camera at 5 fps and forwards selected images to the local model.'),
('Initial operating environment','Prepared, cleared indoor room with a validated position reference. Unfamiliar-room search and outdoor patrol are not established capabilities.')],[1.4,5.4],10)
h('What is delivered and what remains')
p('Delivered: a detailed BOM, proposed electrical connections, parametric OpenSCAD frame source, camera-only ESP32 firmware, a bounded Python camera adapter, a local bench viewer and optional model forwarding. The full assembly and software instructions are in hardware/esp32/README.md.')
p('Remaining: final supplier selection, routed carrier PCB/Gerbers, printed-fit and thrust tests, flight-firmware port, motor calibration, estimator integration and an ESP-Drone vehicle adapter. The supplied camera image holds all motor gates low and cannot fly. No physical camera or aircraft was exercised.')
h('Cost boundary')
p('The $48.50 version supports camera and stabilization development hardware; it does not include demonstrated position hold. The $69 sensor pack supports development of local motion estimation, not global room coordinates or wall detection. A room reference and appropriate obstacle strategy are still needed.')
p('Budgets exclude shipping, tax, labour, PCB/tooling engineering, software, ground computer/network, charger, USB cable, independent RC equipment, room tracking, obstacle sensors and spares. Shared equipment must be accounted for when calculating a complete system cost.')

page('List of pieces')
p('USD totals include the quantities shown. The camera board is listed at $13.90 [1] and the flow sensor at $16.50 [7]. All other amounts are procurement allowances. Exact motor winding, shaft, battery, carrier and guard fabrication remain to be confirmed.')
rows=[]
labels={
'camera':'XIAO ESP32-S3 Sense with camera and antenna',
'imu':'MPU6050 breakout with 3.3 V logic',
'motors':'8520 3.7 V brushed motors, CW/CCW pairs',
'props':'Matched 55 mm propellers',
'frame':'Frame with motor holders and guards',
'battery':'1S 450–600 mAh high-discharge LiPo',
'mosfets':'AO3400A motor-driver MOSFETs',
'diodes':'SS34 flyback diodes',
'gates':'100 ohm gate resistors + 10 kohm pulldowns',
'caps':'Motor and power decoupling capacitors',
'carrier':'Carrier PCB, IMU mount and connectors',
'wire':'Wiring, strap, camera mount and fasteners',
'voltage':'Battery divider and ADC filter'}
for x in BOM['basic_parts']:
    rows.append((labels[x['id']], '4 + 4' if x['id']=='gates' else str(x['quantity']),f"{x['quantity']*x['unit_usd']:.2f}"))
rows.append(('Basic camera aircraft allowance','',f'{basic:.2f}'))
for x in BOM['navigation_additions']:
    rows.append((x['part'],str(x['quantity']),f"{x['quantity']*x['unit_usd']:.2f}"))
rows.append(('Aircraft with navigation sensor pack','',f'{nav:.2f}'))
table(['Piece','Quantity','USD total'],rows,[5.15,.7,.95],9.5)
h('Selection details that affect the build')
p('8520 specifies motor dimensions, not a complete motor specification. Select one matched set with a compatible shaft and propeller hub; verify current, thrust and mass. Use two clockwise and two counterclockwise motor/prop pairs. Drona documents 3.7 V 8520-class motors and small-drone propellers [5]; this proposal has not tested a matched propulsion set.')
p('Seeed camera modules can change between deliveries [2]. Keep the supplied camera and antenna, record the sensor PID, and verify focus at the intended inspection distance. Do not assume the older OV2640 camera is what arrives.')
p('The carrier cost covers material/assembly components only. No routed PCB is supplied. The MOSFET and diode selections require switching, cooling and stall-current validation; their datasheet maximum current is not a tested board current rating.')

page('Mechanical assembly and wiring')
h('Frame and mounting')
p('frame.scad defines a 100 mm motor-diagonal layout, approximately 131 mm square across guards, with nominal 8.6 mm motor bores. It is an unprinted dimensional prototype. Tune the bore to the motor and process; use retention sleeves/rings. Confirm stiffness and guard clearance before treating it as a flight part.')
p('The proposed motor can top is z=8 mm, bottom z=-12 mm, propeller plane z=12 mm, guard height z=8–16 mm and feet z=-14 mm. Actual shaft/hub height controls this geometry. Centre the battery under the body; mount the camera facing forward, antenna clear of wiring and IMU rigidly aligned. Keep the downward sensor view unobstructed.')
p('Component allowances total 67.3 g basic and 71.8 g with sensors. These are not weighed values. Require a proposed 2:1 total maximum thrust-to-weight margin near the low operating voltage, measured with the actual battery and propellers. Flight duration is unknown.')
table(['Connection','XIAO pin and destination'],[
('Four motor gates','M1 front-right CW: D0/GPIO1; M2 rear-right CCW: D1/GPIO2; M3 rear-left CW: D3/GPIO4; M4 front-left CCW: D8/GPIO7.'),
('IMU','SDA: D4/GPIO5; SCL: D5/GPIO6; interrupt: D9/GPIO8. Use 3.3 V logic and verify breakout power/pull-ups.'),
('Battery sensing','D10/GPIO9 through 100k/100k divider with 100 nF filter to ground; calibrate against a meter.'),
('Optional flow sensor','D6/GPIO43 TX → sensor RX; D7/GPIO44 RX ← sensor TX. Separate regulated 5 V sensor supply; 3.3 V UART logic.'),
('Camera','Uses the Sense board’s native camera connector and vendor pin mapping [2]. Keep SD card removed and SD/microphone software disabled.')],[1.2,5.6],9.5)
h('One motor channel repeated four times')
p('Battery positive → motor positive; motor negative → MOSFET drain; MOSFET source → ground. GPIO → 100 ohm → gate; 10 kohm from gate to ground. Place the flyback diode across the motor with cathode to battery positive and anode to motor negative. Fit 100 nF across each motor terminal. Motor current must not flow through an ESP32 GPIO or its regulator.')
p('Split battery power into motor and electronics branches with common ground. The 1S battery feeds the XIAO BAT pads according to its exact board revision, not its 5 V pin. Add a 470 uF >=6.3 V motor-branch capacitor and local 47 uF/100 nF electronics decoupling, with stability checks. The optional sensor’s 5 V boost output must not connect to BAT. Initial USB camera tests use disconnected motors and flight battery.')
p('connections.json is the carrier connection schedule, not a verified PCB. GPIO7/8/9 conflict with SD use. Leave GPIO3, GPIO21 and GPIO41/42 unused by the carrier. Any future flight firmware must explicitly map this motor order and IMU axes.')

page('Live vision into drone137')
h('Camera firmware and receiver')
p('The camera_bench sketch exposes /capture on the local Wi-Fi address. It produces JPEGs with boot ID, frame sequence, device capture time and dimensions. It uses VGA capture, quality 14 and two PSRAM buffers, rejects frames above 200 KB and holds the proposed motor gates LOW. Initialization is not proof of focus or usable flight imagery.')
p('Use the official ESP32 Arduino package pinned to esp32:esp32@3.3.12, board XIAO_ESP32S3 and OPI PSRAM. Copy config.example.h to config.h and set local Wi-Fi credentials. Keep config.h private. The complete commands and actual compile result are recorded in the accompanying README and validation.json.')
p('Compile from the repository root:',None)
p('arduino-cli compile --fqbn esp32:esp32:XIAO_ESP32S3:PSRAM=opi hardware/esp32/camera_bench')
p('After manually flashing the camera-only image, use the device IP printed by its serial console. Start the receiver from the repo root:')
p('.venv/bin/python -m scripts.esp32_camera_bench --camera http://192.168.1.50/capture --seconds 30')
p('Replace the example IP with the device address. Open http://127.0.0.1:8144 while the command runs. The viewer shows the latest camera image and status. The run saves latest.jpg and report.json. To send observations to an already running local model, add:')
p('--model http://127.0.0.1:8081 --inference-every 3 --inference-timeout 30')
h('Data contract and memory behaviour')
p('Esp32Camera implements the existing FrameSource contract. The receiver retains one latest JPEG, samples at 5 fps by default and supports at most one outstanding model request. Selected images use LlamaBackend.image and the existing object-survey schema. Successful model observations retain their source image and description on disk; they do not authorize motion.')
p('The receiver rejects malformed/oversized images, mismatched dimensions, repeated capture clocks, out-of-order frame IDs and HTTP redirects. Network errors clear the current frame. Frames exceeding the host receipt-age limit are unavailable to consumers. Device and host clocks are not synchronized, so capture age is reported as unknown. Stock-camera compatibility mode is explicitly marked as receipt-clock-only.')
h('What the checks establish')
p('All 12 HTTP camera tests passed using fixture images and a model stub. The camera sketch compiled successfully with ESP32 core 3.3.12: 955,754 bytes of program storage and 56,336 bytes of static global memory. These checks do not measure recognition accuracy, runtime memory, physical camera performance, wireless latency or flight. See validation.json for the recorded outcomes.')

page('Flight completion work and sources')
h('Requirements before autonomous room search')
p('1. Finalize the carrier PCB and selected parts, print/fit the frame, verify power under load and measure thrust/temperature. Confirm guards, camera visibility and the actual centre of gravity. Obtain complete delivered prices before claiming the under-$50 target.')
p('2. Port and validate flight firmware. ESP-Drone is a candidate basis [4], not a drop-in image for this carrier. It needs pin/IMU support, motor calibration, mixer mapping, arming and landing logic, battery monitoring and command-loss handling. Camera/Wi-Fi tasks must not violate control deadlines. If this cannot be demonstrated on one MCU, revise the design to separate flight and vision processing and update cost/mass.')
p('3. Integrate positioning. The MTF-02P requires its own UART driver and estimator integration; it is not the same sensor set as the ESP-Drone reference flow deck. Optical flow and downward range do not locate walls or give global room coordinates. Develop a measured-marker or other validated room reference before position-directed missions.')
p('4. Add the ESP-Drone command and telemetry adapter to drone137. The existing physical-control direction is PX4/MAVSDK; this is a different firmware/protocol. First test the adapter and failure cases in simulation. Keep physical_enabled=False until commissioning and project acceptance requirements are met.')
p('5. Commission one aircraft in a prepared area. Test radio loss, stalled inference, stale vision, bad positioning, low battery and operator intervention. Then test object-present and object-absent missions on an actual 8 GB ground computer. No real-flight, 100% recognition, outdoor-weather or multi-drone capability is claimed by this package.')
h('Primary sources checked 22 October 2026')
for title,url,note in [
('[1] Seeed XIAO ESP32-S3 Sense','https://www.seeedstudio.com/XIAO-ESP32S3-Sense-p-5639.html','Camera-board listing, approximately $13.90.'),
('[2] Seeed board and camera reference','https://wiki.seeedstudio.com/xiao_esp32s3_getting_started/','Pin mapping and camera-module variations.'),
('[3] Seeed camera usage','https://wiki.seeedstudio.com/xiao_esp32s3_camera_usage/','Camera bring-up examples and setup.'),
('[4] Espressif ESP-Drone','https://github.com/espressif/esp-drone','Controller basis; limited support and extra sensors for hold modes.'),
('[5] Drona hardware parts','https://www.dronaaviation.com/hardware/parts','Brushed motor and propeller reference.'),
('[6] Alpha and Omega AO3400A','https://www.aosmd.com/products/mosfets/low-voltage-mosfets-12v-30v/ao3400a','Candidate driver electrical limits.'),
('[7] MicoAir MTF-02P store','https://store.micoair.com/product/mtf-02p/','Listed $16.50 sensor.'),
('[8] MicoAir MTF-02P specifications','https://micoair.com/optical_range_sensor_mtf-02p/','Power, serial interface and operating conditions.')]:linkpara(title,url,note)
p('Package files: hardware/esp32/{README.md, bom.json, connections.json, frame.scad, camera_bench/, validation.json}; adapters/esp32_camera.py; scripts/esp32_camera_bench.py; tests/test_esp32_camera.py. The earlier PX4/Pi proposal remains available as a separate design alternative.')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
